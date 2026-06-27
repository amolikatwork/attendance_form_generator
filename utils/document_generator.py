from pathlib import Path
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def create_single_docx(record, path, company_name, month, department=""):
    """
    Create a single Word document for manual entry.
    
    Args:
        record: EmployeeRecord object
        path: Output file path
        company_name: Company name
        month: Month string
        department: Department name (optional)
    """
    document = Document()

    # Set margins
    for section in document.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Title
    title = document.add_heading(company_name, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Form title
    subtitle = document.add_paragraph("Attendance Correction Form")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(14)

    # Employee details
    document.add_paragraph(f"Employee Name: {record.name}")
    document.add_paragraph(f"Employee Code: {record.code}")
    if department:
        document.add_paragraph(f"Department: {department}")
    document.add_paragraph(f"Month: {month}")
    document.add_paragraph(f"Date: {datetime.now().strftime('%d-%m-%Y')}")

    # Attendance table
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    
    # Header
    header_cells = table.rows[0].cells
    header_cells[0].text = "S.No"
    header_cells[1].text = "Date"
    header_cells[2].text = "Reason"
    
    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    # Data rows
    for index, issue in enumerate(record.issues, start=1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(index)
        row_cells[1].text = issue.date
        row_cells[2].text = issue.reason

    # Spacing
    document.add_paragraph("")
    document.add_paragraph("")

    # Signature table
    signature_table = document.add_table(rows=2, cols=2)
    signature_table.style = "Table Grid"
    
    sig_cells = signature_table.rows[0].cells
    sig_cells[0].text = "Employee Signature"
    sig_cells[1].text = "HR Signature"
    
    sig_cells = signature_table.rows[1].cells
    sig_cells[0].text = "_" * 30
    sig_cells[1].text = "_" * 30

    document.save(path)


def convert_to_pdf(docx_path):
    """
    Convert DOCX to PDF.
    Returns path to PDF or None if conversion fails.
    """
    try:
        from docx2pdf import convert
        pdf_path = docx_path.replace(".docx", ".pdf")
        convert(docx_path, pdf_path)
        return pdf_path
    except Exception:
        pass
    
    try:
        import subprocess
        pdf_path = docx_path.replace(".docx", ".pdf")
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(Path(docx_path).parent),
                docx_path,
            ],
            check=True,
            capture_output=True,
            timeout=30,
        )
        return pdf_path
    except Exception:
        pass
    
    return None
